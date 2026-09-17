from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from pathlib import Path

ROOT = Path('/Users/youbin/Desktop/personaAI')
SOURCE_IMG = ROOT / 'tmp/pdfs/source-1.png'
LOGO = ROOT / 'tmp/pdfs/logo.png'
OUT = ROOT / 'output/2026-2학기_주간학습보고서_0주차_워드양식.docx'

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margins(cell, top=40, start=70, bottom=40, end=70):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in('w:tcBorders')
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV', 'start', 'end'):
        if edge in kwargs:
            element = tcBorders.find(qn('w:' + edge))
            if element is None:
                element = OxmlElement('w:' + edge)
                tcBorders.append(element)
            for key in ['val', 'sz', 'space', 'color']:
                if key in kwargs[edge]:
                    element.set(qn('w:' + key), str(kwargs[edge][key]))

def set_table_borders(table, color='000000', size=8):
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell,
                            top={'val': 'single', 'sz': size, 'color': color},
                            bottom={'val': 'single', 'sz': size, 'color': color},
                            start={'val': 'single', 'sz': size, 'color': color},
                            end={'val': 'single', 'sz': size, 'color': color})

def set_width(cell, width_in):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn('w:tcW'))
    if tcW is None:
        tcW = OxmlElement('w:tcW')
        tcPr.append(tcW)
    tcW.set(qn('w:w'), str(int(width_in * 1440)))
    tcW.set(qn('w:type'), 'dxa')

def set_row_height(row, height_in):
    row.height = Inches(height_in)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY

def set_font(run, size=10.5, bold=False, color='000000'):
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=0, line_spacing=1.0):
    p.alignment = align
    fmt = p.paragraph_format
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = line_spacing

def put(cell, text='', size=10.5, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, color='000000'):
    cell.text = ''
    p = cell.paragraphs[0]
    format_paragraph(p, align=align)
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    return p

def add_runs(p, segments, size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    p.clear()
    format_paragraph(p, align=align)
    for text, bold, color in segments:
        r = p.add_run(text)
        set_font(r, size=size, bold=bold, color=color)

im = Image.open(SOURCE_IMG)
im.crop((105, 85, 405, 165)).save(LOGO)

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
sec.top_margin = Inches(0.46)
sec.bottom_margin = Inches(0.45)
sec.left_margin = Inches(0.64)
sec.right_margin = Inches(0.64)
sec.header_distance = Inches(0.2)
sec.footer_distance = Inches(0.2)

style = doc.styles['Normal']
style.font.size = Pt(10.5)

top = doc.add_table(rows=1, cols=2)
top.alignment = WD_TABLE_ALIGNMENT.CENTER
top.autofit = False
set_width(top.cell(0, 0), 5.25)
set_width(top.cell(0, 1), 1.75)
for c in top.rows[0].cells:
    set_cell_border(c, top={'val':'nil'}, bottom={'val':'nil'}, start={'val':'nil'}, end={'val':'nil'})
    set_cell_margins(c, 0, 0, 0, 0)
p = top.cell(0,0).paragraphs[0]
format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT)
p.add_run().add_picture(str(LOGO), width=Inches(1.85))
team = top.cell(0,1).add_table(rows=1, cols=2)
team.alignment = WD_TABLE_ALIGNMENT.RIGHT
team.autofit = False
set_width(team.cell(0,0), 0.70)
set_width(team.cell(0,1), 1.05)
put(team.cell(0,0), '팀번호', size=9.5, bold=True)
set_cell_shading(team.cell(0,0), 'E2F0D9')
put(team.cell(0,1), '50', size=10.5, color='45459B')
set_table_borders(team, size=6)

sp = doc.add_paragraph()
sp.paragraph_format.space_after = Pt(2)

p = doc.add_paragraph()
format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(13))
r = p.add_run('2026-2학기  세종창의학기제  주간학습보고서  ')
set_font(r, size=19, bold=True)
r = p.add_run('(1주차)')
set_font(r, size=19, bold=True, color='0000FF')

LABEL = 'E2F0D9'
GRAY = 'D9D9D9'

info = doc.add_table(rows=5, cols=6)
info.alignment = WD_TABLE_ALIGNMENT.CENTER
info.autofit = False
widths = [1.02, 1.88, 0.84, 1.70, 0.85, 0.71]
for row in info.rows:
    for i, c in enumerate(row.cells):
        set_width(c, widths[i])
        set_cell_margins(c, 30, 45, 30, 45)

set_row_height(info.rows[0], 0.34)
put(info.cell(0,0), '창의과제', size=10, bold=True); set_cell_shading(info.cell(0,0), LABEL)
put(info.cell(0,1).merge(info.cell(0,5)), '페르소나 기반 챗봇 응답 품질 평가 플랫폼 개발', size=10, align=WD_ALIGN_PARAGRAPH.LEFT)

set_row_height(info.rows[1], 0.34)
put(info.cell(1,0), '이름', size=10, bold=True); set_cell_shading(info.cell(1,0), LABEL)
put(info.cell(1,1), '이유빈', size=10)
put(info.cell(1,2), '학습기간', size=10, bold=True); set_cell_shading(info.cell(1,2), LABEL)
put(info.cell(1,3).merge(info.cell(1,5)), '9.3~9.10', size=10)

set_row_height(info.rows[2], 0.34)
put(info.cell(2,0), '학번', size=10, bold=True); set_cell_shading(info.cell(2,0), LABEL)
put(info.cell(2,1), '22011969', size=10)
put(info.cell(2,2), '학습주차', size=10, bold=True); set_cell_shading(info.cell(2,2), LABEL)
put(info.cell(2,3), '1주차', size=10)
put(info.cell(2,4), '학습시간', size=10, bold=True); set_cell_shading(info.cell(2,4), LABEL)
put(info.cell(2,5), '6', size=10)

set_row_height(info.rows[3], 0.40)
put(info.cell(3,0), '학과(전공)', size=10, bold=True); set_cell_shading(info.cell(3,0), LABEL)
put(info.cell(3,1), '지능기전공학부\n스마트기기전공', size=9.5)
put(info.cell(3,2), '과목명', size=10, bold=True); set_cell_shading(info.cell(3,2), LABEL)
put(info.cell(3,3), '자기주도창의전공3', size=9.5)
put(info.cell(3,4), '수강학점', size=10, bold=True); set_cell_shading(info.cell(3,4), LABEL)
put(info.cell(3,5), '3', size=10)

set_row_height(info.rows[4], 0.32)
note = info.cell(4,0).merge(info.cell(4,5))
put(note, '※ 수강학점에 따른 회차별 학습시간 및 10회차 이상 학습 준수', size=9.5, bold=True, color='0000FF')
set_cell_shading(note, GRAY)
set_table_borders(info, size=8)

doc.add_paragraph().paragraph_format.space_after = Pt(2)

form = doc.add_table(rows=6, cols=2)
form.alignment = WD_TABLE_ALIGNMENT.CENTER
form.autofit = False
for row in form.rows:
    set_width(row.cells[0], 1.02)
    set_width(row.cells[1], 6.00)
    for c in row.cells:
        set_cell_margins(c, 45, 55, 45, 55)

heights = [0.42, 2.10, 0.40, 0.68, 0.68, 0.68]
labels = ['금주\n학습목표', '학습내용', '학습방법', '학습성과\n및\n목표달성도', '참고자료\n및 문헌', '내주 계획']
for i, (h, label) in enumerate(zip(heights, labels)):
    set_row_height(form.rows[i], h)
    put(form.cell(i,0), label, size=10.5, bold=True)
    set_cell_shading(form.cell(i,0), LABEL)
    put(form.cell(i,1), '', size=10, align=WD_ALIGN_PARAGRAPH.LEFT)

goal = form.cell(0,1)
goal.text = ''
p = goal.paragraphs[0]
format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT)
add_runs(p, [('personaAI에서  구현해야할  주요 ', False, '000000'), ('기능과 화면을 정의.', True, '000000')], size=10.2)
p2 = goal.add_paragraph()
format_paragraph(p2, align=WD_ALIGN_PARAGRAPH.LEFT)
add_runs(p2, [('사용자 흐름 및 전체 시스템 구현 구조를 설계', True, '000000')], size=10.2)

set_table_borders(form, size=8)

doc.add_paragraph().paragraph_format.space_after = Pt(0)
sig = doc.add_table(rows=2, cols=4)
sig.alignment = WD_TABLE_ALIGNMENT.CENTER
sig.autofit = False
for row in sig.rows:
    for i, c in enumerate(row.cells):
        set_width(c, [3.0, 0.85, 0.65, 2.50][i])
        set_cell_border(c, top={'val':'nil'}, bottom={'val':'nil'}, start={'val':'nil'}, end={'val':'nil'})
        set_cell_margins(c, 0, 0, 0, 0)
set_row_height(sig.rows[0], 0.32)
put(sig.cell(0,0), '', size=10)
put(sig.cell(0,1), '2026년', size=10, bold=True)
put(sig.cell(0,2), '월', size=10, bold=True)
put(sig.cell(0,3), '일', size=10, bold=True)
set_row_height(sig.rows[1], 0.36)
put(sig.cell(1,0), '지도교수', size=10.5, bold=True)
put(sig.cell(1,1), '', size=10)
put(sig.cell(1,2), '', size=10)
put(sig.cell(1,3), '(인)', size=10.5, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)

for table in doc.tables:
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        cantSplit = OxmlElement('w:cantSplit')
        trPr.append(cantSplit)

doc.save(OUT)
print(OUT)
